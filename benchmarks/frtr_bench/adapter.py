import json
import re
import shutil
from collections import defaultdict

import httpx

from bench.benchmark import Benchmark, ChatRequest, Grade, Sample
from bench.common import download, read_jsonl, write_jsonl
from bench.office import xlsx_drop_sheet
from bench.sandbox.exec import extract_code, run_python

REPO = "AnmolGulati6/FRTR-bench"
REVISION = "aed8170d61fed9795b91b1396c60f5b2e5fe0f82"
REF = re.compile(r"'?([A-Za-z0-9 _&()\-.]+?)'?!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")

PROMPT = """You are a financial analyst working with an Excel workbook at:
{path}

Answer the question below. You can inspect the workbook by writing Python (openpyxl and pandas are available) in a single ```python code block per reply; you will see its output. Explore the sheets as needed — values may need lookup, aggregation or reasoning across sheets.

Question: {question}

When you are confident, reply WITHOUT any code block, ending with a line:
FINAL ANSWER: <answer>"""

JUDGE = """You are grading an answer to a spreadsheet question. The benchmark's reference answer is given as a cell reference or expression, resolved below to the actual workbook values.

Question: {question}
Reference answer (as specified by the benchmark): {spec}
Provenance: {provenance}
Resolved values from the workbook:
{resolved}

Model's final answer: {answer}

Decide whether the model's answer is correct (numbers may differ by rounding or formatting; a label answer must name the right item). Reply with JSON: {{"correct": true|false, "reason": "..."}}"""


def _cells(wb, sheet, rng, limit=40):
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    out = []
    cells = ws[rng.replace("$", "")]
    rows = cells if isinstance(cells, tuple) else ((cells,),)
    for row in rows:
        row = row if isinstance(row, tuple) else (row,)
        for c in row:
            if len(out) >= limit:
                return out
            head = ws.cell(row=1, column=c.column).value or ws.cell(row=2, column=c.column).value
            label = ws.cell(row=c.row, column=1).value
            out.append(f"{sheet}!{c.coordinate} = {c.value!r} (column header: {head!r}; row label: {label!r})")
    return out


class FRTR(Benchmark):
    def fetch(self):
        files = httpx.get(f"https://api.github.com/repos/{REPO}/contents?ref={REVISION}", timeout=60).json()
        import openpyxl

        rows = []
        for f in files:
            if not f["name"].endswith(".xlsx"):
                continue
            src = download(f"https://raw.githubusercontent.com/{REPO}/{REVISION}/{f['name']}", self.data_dir / "raw" / f["name"])
            wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
            qname = next((n for n in wb.sheetnames if n.lower().startswith("question")), None)
            if not qname:
                continue
            ws = wb[qname]
            it = ws.iter_rows(values_only=True)
            header = [str(h or "").strip() for h in next(it)]
            for i, r in enumerate(it):
                rec = dict(zip(header, r))
                if not rec.get("Question"):
                    continue
                rows.append({"id": f"{f['name'][:-5]}/{i}", "workbook": f["name"], "question": rec["Question"],
                             "spec": str(rec.get("Answer") or ""), "provenance": str(rec.get("Provenance") or ""),
                             "type": rec.get("ReasoningType"), "difficulty": rec.get("Difficulty")})
            wb.close()
            (self.data_dir / "model").mkdir(exist_ok=True)
            xlsx_drop_sheet(src, self.data_dir / "model" / f["name"], qname)
        write_jsonl(self.data_dir / "questions.jsonl", rows)
        self.mark_fetched()

    def load(self):
        return [Sample(id=r["id"], data=r, gold=r["spec"], stratum=r.get("difficulty")) for r in read_jsonl(self.data_dir / "questions.jsonl")]

    def _ws(self, sample, ctx):
        return ctx.run_dir / "work" / sample.id.replace("/", "__")

    def build(self, sample, ctx):
        ws = self._ws(sample, ctx)
        if ws.exists():
            shutil.rmtree(ws)
        ws.mkdir(parents=True)
        book = ws / sample.data["workbook"]
        shutil.copy2(self.data_dir / "model" / sample.data["workbook"], book)
        max_turns = int(self.spec.options.get("max_turns", 8))
        timeout = float(self.spec.options.get("exec_timeout_s", 120))

        def observe(reply, turn):
            code = extract_code(reply)
            if not code:
                return None
            r = run_python(code, cwd=ws, timeout=timeout)
            res = (r.stdout[-4000:] + ("\n" + r.stderr[-2000:] if r.returncode else "")).strip() or "(no output)"
            if turn + 2 >= max_turns:
                res += "\n\nThis was your last code execution. Reply now with your FINAL ANSWER and no code."
            return f"Execution result:\n{res}"

        return ChatRequest(messages=[{"role": "user", "content": PROMPT.format(path=book, question=sample.data["question"])}],
                           observe=observe, max_turns=max_turns)

    def _resolved(self, sample):
        import openpyxl

        wb = openpyxl.load_workbook(self.data_dir / "raw" / sample.data["workbook"], data_only=True)
        lines = []
        for sheet, rng in REF.findall(sample.data["spec"] + " ; " + sample.data["provenance"]):
            lines += _cells(wb, sheet.strip(), rng)
        return "\n".join(dict.fromkeys(lines)) or "(no cell references could be resolved)"

    def grade(self, sample, output, ctx):
        m = re.findall(r"FINAL ANSWER\s*:\s*(.+)", output.text or "", re.I)
        answer = m[-1].strip() if m else (output.text or "").strip()[-300:]
        if not answer:
            return Grade(score=0.0, details={"error": "no answer"})
        judge = ctx.require_judge()
        verdict, info = judge.score_json(JUDGE.format(question=sample.data["question"], spec=sample.data["spec"],
                                                      provenance=sample.data["provenance"], resolved=self._resolved(sample),
                                                      answer=answer), "rubric")
        return Grade(score=float(bool(verdict.get("correct"))), extracted=answer, judge={**info, "verdict": verdict},
                     details={"difficulty": sample.data.get("difficulty"), "type": sample.data.get("type")})

    def aggregate(self, rows):
        by = defaultdict(list)
        for r in rows:
            by[(r["details"] or {}).get("difficulty")].append(r["score"])
        return {"by_difficulty": {k: sum(v) / len(v) for k, v in by.items() if k}}


BENCHMARK = FRTR
