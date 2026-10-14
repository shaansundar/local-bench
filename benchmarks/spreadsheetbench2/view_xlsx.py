"""Inspect an .xlsx file (stands in for SpreadsheetBench 2's `view_xlsx` tool).

  python3 view_xlsx.py FILE list
  python3 view_xlsx.py FILE content "Sheet1" [START_ROW END_ROW]
"""
import sys

import openpyxl


def main():
    path, mode = sys.argv[1], sys.argv[2]
    wf = openpyxl.load_workbook(path, data_only=False)
    if mode == "list":
        for ws in wf.worksheets:
            print(f"{ws.title}\tdims={ws.dimensions}\tmax_row={ws.max_row}\tmax_col={ws.max_column}")
        return
    sheet = sys.argv[3]
    start = int(sys.argv[4]) if len(sys.argv) > 4 else 1
    end = int(sys.argv[5]) if len(sys.argv) > 5 else min(start + 49, wf[sheet].max_row)
    wv = openpyxl.load_workbook(path, data_only=True)
    ws_f, ws_v = wf[sheet], wv[sheet]
    for r in range(start, end + 1):
        cells = []
        for c in range(1, ws_f.max_column + 1):
            f = ws_f.cell(r, c).value
            if f is None:
                continue
            v = ws_v.cell(r, c).value
            ref = ws_f.cell(r, c).coordinate
            cells.append(f"{ref}={f!r}" + (f" -> {v!r}" if isinstance(f, str) and f.startswith("=") else ""))
        if cells:
            print(f"row {r}: " + " | ".join(cells))


if __name__ == "__main__":
    main()
