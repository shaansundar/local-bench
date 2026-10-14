"""Office-file helpers: LibreOffice recalculation/rendering and workbook/presentation text dumps."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import threading
from pathlib import Path

from bench.config import load_config

_lo_lock = threading.Lock()  # one soffice at a time per profile


def find_soffice() -> str | None:
    for c in ("soffice", "libreoffice"):
        p = shutil.which(c)
        if p:
            return p
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


class LibreOfficeMissing(RuntimeError):
    pass


def _convert(src: Path, fmt: str, outdir: Path, timeout: float = 180) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise LibreOfficeMissing("LibreOffice not found (brew install --cask libreoffice)")
    profile = load_config().results_dir / "tmp" / "lo_profile"
    profile.mkdir(parents=True, exist_ok=True)
    with _lo_lock:
        r = subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                            "--convert-to", fmt, "--outdir", str(outdir), str(src)],
                           capture_output=True, text=True, timeout=timeout)
    ext = fmt.split(":")[0]
    out = outdir / f"{src.stem}.{ext}"
    if r.returncode != 0 or not out.exists():
        raise RuntimeError(f"LibreOffice conversion failed: {r.stderr[-400:]}")
    return out


def recalc_xlsx(src: Path, dest: Path | None = None) -> Path:
    """Open + save through LibreOffice so formula results are cached (openpyxl data_only=True)."""
    dest = dest or src
    with tempfile.TemporaryDirectory() as tmp:
        out = _convert(src, "xlsx:Calc MS Excel 2007 XML", Path(tmp))
        shutil.move(str(out), dest)
    return dest


def to_pdf(src: Path, outdir: Path) -> Path:
    return _convert(src, "pdf", outdir, timeout=300)


def workbook_text(path: Path, max_rows: int = 30, max_cols: int = 30) -> str:
    """Compact text preview of every sheet (values, first rows) for prompts."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet '{ws.title}' (dims {ws.dimensions})")
        for i, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True)):
            if any(v is not None for v in row):
                parts.append(f"{i + 1}: " + " | ".join("" if v is None else str(v)[:40] for v in row))
    wb.close()
    return "\n".join(parts)


def has_formulas(path: Path) -> bool:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    return True
        return False
    finally:
        wb.close()


def needs_recalc(path: Path) -> bool:
    """True when some formula cell has no cached value (e.g. the file was written by openpyxl)."""
    import openpyxl

    if not has_formulas(path):
        return False
    wf = openpyxl.load_workbook(path, read_only=True)
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sf, sv in zip(wf.worksheets, wv.worksheets):
            for rf, rv in zip(sf.iter_rows(values_only=True), sv.iter_rows(values_only=True)):
                for f, v in zip(rf, rv):
                    if isinstance(f, str) and f.startswith("=") and v is None:
                        return True
        return False
    finally:
        wf.close()
        wv.close()


def recalc_if_needed(src: Path, dest: Path) -> Path:
    """Recalculate through LibreOffice when formula results are not cached; otherwise copy as-is."""
    if needs_recalc(src):
        return recalc_xlsx(src, dest)
    shutil.copy2(src, dest)
    return dest


def xlsx_drop_sheet(src: Path, dest: Path, sheet_name: str) -> None:
    """Remove one worksheet at the XML level, preserving everything else (images, charts, styles) —
    openpyxl would silently drop drawings on save."""
    import re
    import zipfile

    with zipfile.ZipFile(src) as zin:
        wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
        rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        m = re.search(r'<sheet\b[^>]*\bname="' + re.escape(sheet_name) + r'"[^>]*/>', wb_xml)
        if not m:
            shutil.copy2(src, dest)
            return
        rid = re.search(r'r:id="([^"]+)"', m.group(0)).group(1)
        rel = re.search(r'<Relationship\b[^>]*\bId="' + re.escape(rid) + r'"[^>]*/>', rels)
        target = re.search(r'Target="([^"]+)"', rel.group(0)).group(1).lstrip("/")
        part = target if target.startswith("xl/") else f"xl/{target}"
        part_rels = part.rsplit("/", 1)[0] + "/_rels/" + part.rsplit("/", 1)[1] + ".rels"
        wb_xml = wb_xml.replace(m.group(0), "")
        wb_xml = re.sub(r'<definedName\b[^>]*>[^<]*' + re.escape(sheet_name) + r'[^<]*</definedName>', "", wb_xml)
        rels = rels.replace(rel.group(0), "")
        ctypes = zin.read("[Content_Types].xml").decode("utf-8")
        ctypes = re.sub(r'<Override\b[^>]*PartName="/' + re.escape(part) + r'"[^>]*/>', "", ctypes)
        with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in (part, part_rels):
                    continue
                data = zin.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    data = wb_xml.encode()
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    data = rels.encode()
                elif item.filename == "[Content_Types].xml":
                    data = ctypes.encode()
                zout.writestr(item, data)


def document_text(path: Path, max_chars: int = 2_000_000) -> str:
    """Plain-text rendering of an office/PDF/HTML/CSV document for prompts."""
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for ti, t in enumerate(d.tables):
            parts.append(f"[Table {ti + 1}]")
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("[Notes] " + slide.notes_slide.notes_text_frame.text)
        text = "\n".join(parts)
    elif ext in (".xlsx", ".xlsm"):
        text = workbook_text(path, max_rows=2000, max_cols=60)
    elif ext == ".pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        text = "\n\n".join(f"--- Page {i + 1} ---\n" + pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
    elif ext in (".htm", ".html"):
        import trafilatura

        raw = path.read_text(errors="replace")
        text = trafilatura.extract(raw, include_tables=True) or ""
    else:
        text = path.read_text(errors="replace")
    return text[:max_chars]


def render_slides(pptx: Path, outdir: Path, max_slides: int = 20, scale: float = 1.0) -> list[Path]:
    """PPTX -> PDF (LibreOffice) -> PNG per slide (pypdfium2, no poppler needed)."""
    import pypdfium2 as pdfium

    outdir.mkdir(parents=True, exist_ok=True)
    pdf = to_pdf(pptx, outdir)
    doc = pdfium.PdfDocument(str(pdf))
    paths = []
    for i in range(min(len(doc), max_slides)):
        img = doc[i].render(scale=scale).to_pil()
        p = outdir / f"{pptx.stem}_{i + 1:02d}.png"
        img.save(p)
        paths.append(p)
    return paths
